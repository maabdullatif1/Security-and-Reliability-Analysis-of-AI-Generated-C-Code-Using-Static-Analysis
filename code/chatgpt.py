from openai import OpenAI
import time
import os
import csv
import subprocess
from datetime import datetime
import xml.etree.ElementTree as ET
from collections import Counter
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
import shutil
import re


# -------------------------------------------------
# Extract only C++ code (prevents explanation text + ```cpp in saved .cpp)
# (This is the function you used before, with a tiny fallback cleanup.)
# -------------------------------------------------
def extract_cpp_code(response: str) -> str:
    """Extracts C++ code from fenced blocks; otherwise trims explanation text safely."""
    if not response:
        return ""

    # Prefer fenced code blocks if present
    matches = re.findall(r"```(?:cpp|c\+\+)?\n(.*?)\n```", response, re.DOTALL)
    if matches:
        return "\n".join(matches).strip()

    # No fenced block: remove any leftover fence markers
    text = re.sub(r"^```.*?$", "", response, flags=re.MULTILINE).strip()

    # SAFE TRIM: cut off trailing non-code text after the last closing brace
    last_brace = text.rfind("}")
    if last_brace != -1:
        return text[:last_brace + 1].strip()

    # If no braces exist (rare), fallback to last semicolon
    last_semi = text.rfind(";")
    if last_semi != -1:
        return text[:last_semi + 1].strip()

    return text



# -------------------------------------------------
# macOS SDK path helper (fixes 'iostream' file not found for clang-tidy on macOS)
# -------------------------------------------------
def get_macos_sdk_path() -> str:
    try:
        r = subprocess.run(
            ["xcrun", "--sdk", "macosx", "--show-sdk-path"],
            stdout=subprocess.PIPE,
            stderr=subprocess.DEVNULL,
            text=True,
            check=False
        )
        return (r.stdout or "").strip()
    except Exception:
        return ""


def parse_cppcheck_xml(xml_path: str) -> Counter:
    counts = Counter()
    try:
        tree = ET.parse(xml_path)
        root = tree.getroot()
    except Exception:
        return counts

    errors_node = root.find("errors")
    if errors_node is None:
        errors_node = root.find(".//errors")
    if errors_node is None:
        return counts

    for err in errors_node.findall("error"):
        severity = (err.get("severity") or "unknown").strip()
        err_id = (err.get("id") or "unknown").strip()
        counts[(severity, err_id)] += 1

    return counts


def autosize_columns(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = "" if cell.value is None else str(cell.value)
            if len(val) > max_len:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = min(max_len + 2, 70)


def write_table(ws, headers, rows):
    ws.append(headers)
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    autosize_columns(ws)


def find_clang_tidy() -> str:
    # Try PATH first
    p = shutil.which("clang-tidy")
    if p:
        return p

    # Common Homebrew locations
    candidates = [
        "/opt/homebrew/opt/llvm/bin/clang-tidy",   # Apple Silicon
        "/usr/local/opt/llvm/bin/clang-tidy",      # Intel
        "/opt/homebrew/bin/clang-tidy",
        "/usr/local/bin/clang-tidy",
    ]
    for c in candidates:
        if os.path.isfile(c) and os.access(c, os.X_OK):
            return c

    return ""


def ensure_cmake_lists(task_dir: str):
    cmake_path = os.path.join(task_dir, "CMakeLists.txt")
    if os.path.exists(cmake_path):
        return

    cmake_content = """cmake_minimum_required(VERSION 3.16)
project(AIGenerated CXX)
set(CMAKE_CXX_STANDARD 17)
set(CMAKE_CXX_STANDARD_REQUIRED ON)

file(GLOB SRC_FILES "*.cpp")
foreach(src ${SRC_FILES})
  get_filename_component(name ${src} NAME_WE)
  add_executable(${name} ${src})
endforeach()
"""
    with open(cmake_path, "w", encoding="utf-8") as f:
        f.write(cmake_content)


def run_cmake_export_compile_commands(task_dir: str) -> str:
    """
    Creates build/compile_commands.json and returns build directory path.
    Forces Apple clang/clang++ and sets macOS SDK so standard headers resolve.
    """
    ensure_cmake_lists(task_dir)
    build_dir = os.path.join(task_dir, "build")
    os.makedirs(build_dir, exist_ok=True)

    sdk = get_macos_sdk_path()

    cmd = [
        "cmake",
        "-S", task_dir,
        "-B", build_dir,
        "-DCMAKE_EXPORT_COMPILE_COMMANDS=ON",
        "-DCMAKE_C_COMPILER=clang",
        "-DCMAKE_CXX_COMPILER=clang++",
    ]
    if sdk:
        cmd.append(f"-DCMAKE_OSX_SYSROOT={sdk}")

    subprocess.run(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, check=False)
    return build_dir


# More tolerant: captures clang-diagnostic-error too
_CLANG_TIDY_LINE = re.compile(r"^.+?:\d+:\d+:\s+(warning|error|note):.*\[(.+?)\]\s*$")


def parse_clang_tidy_output(text: str) -> Counter:
    """
    Returns counts[(severity, check_name)] += 1
    """
    counts = Counter()
    if not text:
        return counts

    for line in text.splitlines():
        m = _CLANG_TIDY_LINE.match(line.strip())
        if not m:
            continue
        severity = m.group(1).strip()
        check_name = m.group(2).strip() if m.group(2) else "unknown"
        counts[(severity, check_name)] += 1

    return counts


def run_clang_tidy_for_task(task_dir: str, clang_tidy_path: str, checks: str) -> tuple[Counter, list]:
    """
    Runs clang-tidy on each .cpp file in task_dir and returns:
      - aggregated Counter[(severity, check)] counts
      - per-file rows: (file_name, status, diag_count)
    """
    build_dir = run_cmake_export_compile_commands(task_dir)
    cpp_files = sorted([f for f in os.listdir(task_dir) if f.lower().endswith(".cpp")])

    aggregated = Counter()
    per_file_rows = []

    sdk = get_macos_sdk_path()

    for f in cpp_files:
        file_path = os.path.join(task_dir, f)

        cmd = [
            clang_tidy_path,
            file_path,
            "-p", build_dir,
            f"--checks={checks}",
            "--extra-arg=-std=c++17",
        ]

        # Key fix for macOS: allow clang-tidy to find libc++ headers (iostream, etc.)
        if sdk:
            cmd.append("--extra-arg=-isysroot")
            cmd.append(f"--extra-arg={sdk}")

        res = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, check=False)
        out = res.stdout or ""

        # Save raw report per file (auditable)
        report_path = os.path.join(task_dir, f"{os.path.splitext(f)[0]}_clang_tidy.txt")
        with open(report_path, "w", encoding="utf-8") as rpt:
            rpt.write(out)

        counts = parse_clang_tidy_output(out)
        aggregated.update(counts)

        diag_count = sum(counts.values())
        status = "ok" if res.returncode == 0 else "nonzero_exit"
        per_file_rows.append((f, status, diag_count))

    return aggregated, per_file_rows


# -----------------------------
# Configuration
# -----------------------------
MODEL = "gpt-5.2-chat-latest"
RUNS_PER_TASK = 10
TASKS_FILE = "tasks.csv"  # required columns: category, task_id, task_description (FULL PROMPT TEXT)
BASE_DIR = "experiments"
EXEC_DATE = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")

EXCEL_OUT = os.path.join(BASE_DIR, "static_analysis_summary.xlsx")

# clang-tidy checks (security/reliability focused, avoids noisy style-only rules)
CLANG_TIDY_CHECKS = "clang-analyzer-*,bugprone-*,performance-*"

client = OpenAI()  # API key from environment only (no key in code)

# -----------------------------
# Load tasks
# -----------------------------
tasks = []
with open(TASKS_FILE, "r", encoding="utf-8-sig", newline="") as f:
    reader = csv.DictReader(f)
    required = {"category", "task_id", "task_description"}
    if not required.issubset(set(reader.fieldnames or [])):
        raise ValueError(f"{TASKS_FILE} must contain columns: {sorted(required)}")

    for row in reader:
        category = (row.get("category") or "").strip()
        task_id = (row.get("task_id") or "").strip()
        query = (row.get("task_description") or "").strip()  # FULL QUERY TEXT from CSV
        if category and task_id and query:
            tasks.append((category, task_id, query))

if not tasks:
    raise ValueError("No valid tasks found in tasks.csv")

print(f"[INFO] Loaded {len(tasks)} tasks from {TASKS_FILE}", flush=True)

# Find clang-tidy once
CLANG_TIDY_PATH = find_clang_tidy()
if not CLANG_TIDY_PATH:
    print("[WARN] clang-tidy not found. The script will still run Cppcheck and Excel summaries, but clang-tidy will be skipped.", flush=True)
else:
    print(f"[INFO] clang-tidy found at: {CLANG_TIDY_PATH}", flush=True)

sdk_check = get_macos_sdk_path()
if sdk_check:
    print(f"[INFO] macOS SDK: {sdk_check}", flush=True)
else:
    print("[WARN] Could not determine macOS SDK path via xcrun. If clang-tidy still can't find <iostream>, run: xcode-select --install", flush=True)

# -----------------------------
# Generate code + run cppcheck + run clang-tidy
# -----------------------------
cppcheck_task_counters = {}
clang_task_counters = {}
clang_per_file_logs = {}  # (category, task_id) -> list[(file, status, diags)]

for task_index, (category, task_id, query) in enumerate(tasks, start=1):
    category_dir = os.path.join(BASE_DIR, category)
    task_dir = os.path.join(category_dir, task_id)
    os.makedirs(task_dir, exist_ok=True)

    print(f"\n[TASK {task_index}/{len(tasks)}] {category} | {task_id}", flush=True)

    # ---- metadata CSV per task (generation) ----
    metadata_csv = os.path.join(task_dir, f"{task_id}_metadata.csv")
    with open(metadata_csv, "w", newline="", encoding="utf-8") as csvfile:
        writer = csv.writer(csvfile)
        writer.writerow([
            "category",
            "task_id",
            "model",
            "run_id",
            "file_name",
            "execution_date",
            "generation_time_seconds",
            "lines_of_code",
            "character_count"
        ])

        for run in range(1, RUNS_PER_TASK + 1):
            print(f"  [GEN] {task_id} Run {run}/{RUNS_PER_TASK} ...", end="", flush=True)

            start_time = time.time()

            response = client.chat.completions.create(
                model=MODEL,
                messages=[{"role": "user", "content": query}]
            )

            raw_output = response.choices[0].message.content or ""

            # IMPORTANT: prevents explanation text + ```cpp in saved .cpp
            cpp_code = extract_cpp_code(raw_output)

            char_count = len(cpp_code)
            loc = sum(1 for line in cpp_code.splitlines() if line.strip())

            file_name = f"{task_id}_Run{run}.cpp"
            file_path = os.path.join(task_dir, file_name)

            with open(file_path, "w", encoding="utf-8") as out:
                out.write(cpp_code)

            elapsed = time.time() - start_time

            writer.writerow([
                category,
                task_id,
                MODEL,
                run,
                file_name,
                EXEC_DATE,
                round(elapsed, 4),
                loc,
                char_count
            ])

            print(f" done ({elapsed:.2f}s, LOC={loc}, chars={char_count})", flush=True)

    # ---- cppcheck XML named by task_id ----
    print(f"  [CPPCHECK] Running cppcheck for {task_id} ...", flush=True)
    xml_path = os.path.join(task_dir, f"{task_id}.xml")
    cppcheck_cmd = [
        "cppcheck",
        "--enable=all",
        "--xml",
        "--xml-version=2",
        task_dir
    ]
    with open(xml_path, "w", encoding="utf-8") as xml_out:
        subprocess.run(
            cppcheck_cmd,
            stderr=xml_out,               # cppcheck XML is emitted on stderr
            stdout=subprocess.DEVNULL,
            check=False
        )

    cppcheck_counts = parse_cppcheck_xml(xml_path)
    cppcheck_task_counters[(category, task_id)] = cppcheck_counts
    print(f"  [CPPCHECK] Saved {task_id}.xml (unique issues={len(cppcheck_counts)})", flush=True)

    # ---- clang-tidy (automatic: cmake export + analyze) ----
    if CLANG_TIDY_PATH:
        print(f"  [CLANG-TIDY] Running clang-tidy for {task_id} ...", flush=True)

        # Optional but recommended: ensure old build does not carry stale configuration
        # (You can comment this out if you prefer to keep build artifacts.)
        # build_dir = os.path.join(task_dir, "build")
        # if os.path.isdir(build_dir):
        #     shutil.rmtree(build_dir, ignore_errors=True)

        clang_counts, per_file = run_clang_tidy_for_task(
            task_dir=task_dir,
            clang_tidy_path=CLANG_TIDY_PATH,
            checks=CLANG_TIDY_CHECKS
        )
        clang_task_counters[(category, task_id)] = clang_counts
        clang_per_file_logs[(category, task_id)] = per_file

        # Save per-task clang-tidy summary CSV
        clang_csv = os.path.join(task_dir, f"{task_id}_clang_tidy_summary.csv")
        rows = sorted([(sev, chk, cnt) for (sev, chk), cnt in clang_counts.items()],
                      key=lambda x: (-x[2], x[0], x[1]))
        with open(clang_csv, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(["severity", "check_name", "occurrences"])
            for r in rows:
                w.writerow(r)

        # Save per-task clang-tidy per-file status CSV
        per_file_csv = os.path.join(task_dir, f"{task_id}_clang_tidy_files.csv")
        with open(per_file_csv, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(["file_name", "status", "diagnostic_count"])
            for r in per_file:
                w.writerow(r)

        print(f"  [CLANG-TIDY] Saved summary: {task_id}_clang_tidy_summary.csv (unique checks={len(clang_counts)})", flush=True)
        print(f"  [CLANG-TIDY] Saved per-file status: {task_id}_clang_tidy_files.csv", flush=True)
    else:
        clang_task_counters[(category, task_id)] = Counter()
        clang_per_file_logs[(category, task_id)] = []
        print("  [CLANG-TIDY] Skipped (clang-tidy not found).", flush=True)

# -----------------------------
# Build Excel summary (Cppcheck + Clang-Tidy)
# -----------------------------
print("\n[EXCEL] Building Excel summary ...", flush=True)
os.makedirs(BASE_DIR, exist_ok=True)

wb = Workbook()
wb.remove(wb.active)

# ---- Cppcheck summaries ----
cpp_overall = Counter()
cpp_per_task_rows = []

for (category, task_id), counts in cppcheck_task_counters.items():
    for (severity, err_id), cnt in counts.items():
        cpp_overall[(severity, err_id)] += cnt
        cpp_per_task_rows.append((category, task_id, severity, err_id, cnt))

ws_cpp_overall = wb.create_sheet("Cppcheck_Overall")
cpp_overall_rows = sorted([(sev, eid, cnt) for (sev, eid), cnt in cpp_overall.items()],
                          key=lambda x: (-x[2], x[0], x[1]))
write_table(ws_cpp_overall, ["severity", "error_id", "occurrences"], cpp_overall_rows)

ws_cpp_task = wb.create_sheet("Cppcheck_Per_Task")
cpp_per_task_rows_sorted = sorted(cpp_per_task_rows, key=lambda x: (x[0], x[1], -x[4], x[2], x[3]))
write_table(ws_cpp_task, ["category", "task_id", "severity", "error_id", "occurrences"], cpp_per_task_rows_sorted)

# Optional: one sheet per task for cppcheck
for (category, task_id), counts in cppcheck_task_counters.items():
    ws = wb.create_sheet(f"{task_id}_cpp")
    rows = sorted([(sev, eid, cnt) for (sev, eid), cnt in counts.items()],
                  key=lambda x: (-x[2], x[0], x[1]))
    write_table(ws, ["severity", "error_id", "occurrences"], rows)

# ---- Clang-Tidy summaries ----
clang_overall = Counter()
clang_per_task_rows = []

for (category, task_id), counts in clang_task_counters.items():
    for (severity, check_name), cnt in counts.items():
        clang_overall[(severity, check_name)] += cnt
        clang_per_task_rows.append((category, task_id, severity, check_name, cnt))

ws_ct_overall = wb.create_sheet("ClangTidy_Overall")
ct_overall_rows = sorted([(sev, chk, cnt) for (sev, chk), cnt in clang_overall.items()],
                         key=lambda x: (-x[2], x[0], x[1]))
write_table(ws_ct_overall, ["severity", "check_name", "occurrences"], ct_overall_rows)

ws_ct_task = wb.create_sheet("ClangTidy_Per_Task")
ct_per_task_sorted = sorted(clang_per_task_rows, key=lambda x: (x[0], x[1], -x[4], x[2], x[3]))
write_table(ws_ct_task, ["category", "task_id", "severity", "check_name", "occurrences"], ct_per_task_sorted)

# Optional: one sheet per task for clang-tidy
for (category, task_id), counts in clang_task_counters.items():
    ws = wb.create_sheet(f"{task_id}_ct")
    rows = sorted([(sev, chk, cnt) for (sev, chk), cnt in counts.items()],
                  key=lambda x: (-x[2], x[0], x[1]))
    write_table(ws, ["severity", "check_name", "occurrences"], rows)

# Add per-file clang-tidy status (useful if some files fail)
ws_ct_files = wb.create_sheet("ClangTidy_Files")
file_rows = []
for (category, task_id), rows in clang_per_file_logs.items():
    for (fname, status, diag_count) in rows:
        file_rows.append((category, task_id, fname, status, diag_count))
file_rows_sorted = sorted(file_rows, key=lambda x: (x[0], x[1], x[2]))
write_table(ws_ct_files, ["category", "task_id", "file_name", "status", "diagnostic_count"], file_rows_sorted)

wb.save(EXCEL_OUT)
print(f"[DONE] Excel summary saved: {EXCEL_OUT}", flush=True)

